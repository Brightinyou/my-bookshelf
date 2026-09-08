import io
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import zlib

from services.office_extract import extract_docx, extract_hwpx, extract_hwp, _hwp_records, _hwp_paragraph


class OfficeExtractionTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def archive(self, name, parts):
        path = self.root / name
        with zipfile.ZipFile(path, "w") as out:
            for key, value in parts.items():
                out.writestr(key, value)
        return path

    def test_hwpx_real_textbox_and_repeated_text(self):
        import hwpx
        doc = hwpx.HwpxDocument.new()
        p = doc.add_paragraph("본문문장")
        p.add_rectangle().set_draw_text("상자내부문장")
        doc.add_paragraph("본문문장")
        path = self.root / "box.hwpx"
        path.write_bytes(doc.to_bytes())
        result = extract_hwpx(path)
        self.assertIn("상자내부문장", result.text)
        self.assertEqual(result.text.count("본문문장"), 2)
        self.assertEqual(result.textboxes, 1)
        self.assertEqual(result.images, 0)

    def test_hwpx_nested_table_and_section_numeric_order(self):
        section = '<s xmlns:hp="urn:test"><hp:p><hp:run><hp:t>{}</hp:t></hp:run></hp:p></s>'
        nested = '<s xmlns:hp="urn:test"><hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:p><hp:run><hp:t>셀</hp:t><hp:rect><hp:drawText><hp:subList><hp:p><hp:run><hp:t>상자</hp:t></hp:run></hp:p></hp:subList></hp:drawText></hp:rect></hp:run></hp:p></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></s>'
        path = self.archive("nested.hwpx", {"Contents/section10.xml": section.format("마지막"),
                          "Contents/section2.xml": section.format("중간"), "Contents/section0.xml": nested})
        result = extract_hwpx(path)
        self.assertLess(result.text.index("셀"), result.text.index("상자"))
        self.assertLess(result.text.index("중간"), result.text.index("마지막"))
        self.assertEqual(result.text.count("상자"), 1)
        self.assertEqual(result.tables, 1)

    def test_docx_table_real_document(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("시작")
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "표내용"
        doc.add_paragraph("끝")
        path = self.root / "table.docx"
        doc.save(path)
        result = extract_docx(path)
        self.assertLess(result.text.index("시작"), result.text.index("표내용"))
        self.assertLess(result.text.index("표내용"), result.text.index("끝"))

    def test_docx_alternate_textbox_is_not_duplicated_and_notes_survive(self):
        xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="urn:mc"><w:body><w:p><w:r><w:t>본문</w:t><w:footnoteReference w:id="1"/></w:r><mc:AlternateContent><mc:Choice><w:txbxContent><w:p><w:r><w:t>상자</w:t></w:r></w:p></w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent><w:p><w:r><w:t>상자</w:t></w:r></w:p></w:txbxContent></mc:Fallback></mc:AlternateContent><w:del><w:r><w:delText>삭제됨</w:delText></w:r></w:del></w:p></w:body></w:document>'''
        notes = '''<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>각주내용</w:t></w:r></w:p></w:footnote></w:footnotes>'''
        path = self.archive("box.docx", {"word/document.xml": xml, "word/footnotes.xml": notes})
        result = extract_docx(path)
        self.assertEqual(result.text.count("상자"), 1)
        self.assertIn("각주내용", result.text)
        self.assertNotIn("삭제됨", result.text)

    def test_picture_warning_is_visible(self):
        path = self.archive("pic.hwpx", {"Contents/section0.xml": '<s><p><t>본문</t><pic/></p></s>'})
        self.assertIn("OCR", extract_hwpx(path).note)

    def test_hwp_nested_records_and_controls(self):
        def record(text, level=1):
            data = text.encode("utf-16-le")
            return struct.pack("<I", 67 | (level << 10) | (len(data) << 20)) + data
        data = record("본문\r") + record("글상자\r", 4) + record("표내용\r", 7)
        compressor = zlib.compressobj(wbits=-15)
        compressed = compressor.compress(data) + compressor.flush()
        header = b"HWP Document File".ljust(32, b"\0") + bytes([0, 0, 0, 5]) + struct.pack("<I", 1)
        class Ole:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def exists(self, name): return True
            def listdir(self): return [["BodyText", "Section0"]]
            def openstream(self, name): return io.BytesIO(header if name == "FileHeader" else compressed)
        with patch("olefile.OleFileIO", return_value=Ole()):
            result = extract_hwp(self.root / "nested.hwp")
        self.assertEqual(result.text, "본문\n글상자\n표내용")
        tab = b"\x09\x00" + b"x" * 14
        self.assertEqual(_hwp_paragraph("가".encode("utf-16-le") + tab + "나".encode("utf-16-le")), "가\t나")

    def test_truncated_hwp_is_not_silently_accepted(self):
        with self.assertRaises(ValueError):
            list(_hwp_records(struct.pack("<I", 67 | (100 << 20)) + b"x"))
        with self.assertRaises(ValueError):
            _hwp_paragraph(b"\x09\x00")


if __name__ == "__main__":
    unittest.main()
